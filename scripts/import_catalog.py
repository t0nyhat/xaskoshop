#!/usr/bin/env python3
"""Build the storefront catalogue from the admin workbook.

Reads ``catalog/Каталог ХаскоШоп.xlsx`` (sheets «Товары», «Фасовки», «Марки»),
checks it and writes ``assets/data/catalog.js``. Acari products linked to a
manufacturer card get composition, nutrition, photos and feeding tables from
``assets/data/acari-catalog.json`` wherever the admin left the cell empty.

Columns are found by their header, so the admin may reorder them or add own
columns. Errors stop the build and list the rows to fix.

Usage:  python3 -m pip install openpyxl && python3 scripts/import_catalog.py
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "catalog" / "Каталог ХаскоШоп.xlsx"
MAKER = ROOT / "assets/data/acari-catalog.json"
OUTPUT = ROOT / "assets/data/catalog.js"
PAGE = ROOT / "index.html"   # ссылка на catalog.js получает ?v=<хэш данных>, чтобы браузеры не держали старую копию

PRODUCT_COLUMNS = ["ID", "Показывать", "Марка", "Линейка", "Название", "Питомец", "Раздел", "Коротко", "Описание",
                   "Состав", "Белок, %", "Жир, %", "Класс", "Особенности", "Фото", "Карточка производителя",
                   "Порядок", "Комментарий"]
SIZE_COLUMNS = ["ID товара", "Фасовка", "Гранула", "Цена, ₽", "Цена до, ₽", "Остаток", "Показывать", "Штрихкод",
                "Код Saby", "Источник цены", "Комментарий"]
BRAND_COLUMNS = ["ID", "Название", "Подпись", "Цвет", "Класс", "Метка", "Описание", "Фото", "В блоке «Марки»", "Порядок"]
REQUIRED = {"Товары": ["ID", "Марка", "Название", "Питомец", "Раздел"], "Фасовки": ["ID товара", "Фасовка"],
            "Марки": ["ID", "Название"]}

PETS = {"собаки": "dog", "кошки": "cat", "собаки и кошки": "both"}
SECTIONS = ["Корм сухой", "Корм влажный", "Лакомства", "Добавки", "Наполнители", "Аксессуары"]
FEED_TYPE = {"Корм сухой": "сухой", "Корм влажный": "влажный", "Лакомства": "лакомство", "Добавки": "добавка",
             "Наполнители": "наполнитель", "Аксессуары": "аксессуар"}
GRANULES = {"мелкие": "S", "средние": "M", "крупные": "L", "декоративные": "XS"}
YES, NO = {"да", "yes", "1", "+", "true"}, {"нет", "no", "0", "-", "false"}
PRICE_SOURCES = ["ЭЛ каталог", "прайс поставщика", "вручную", "нет цены"]

errors: list[str] = []
warnings: list[str] = []


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def read_sheet(wb, name: str, columns: list[str]) -> list[tuple[int, dict]]:
    if name not in wb.sheetnames:
        errors.append(f"Нет листа «{name}»")
        return []
    rows = list(wb[name].iter_rows(values_only=True))
    header = [text(h) for h in rows[0]] if rows else []
    missing = [c for c in REQUIRED[name] if c not in header]
    if missing:
        errors.append(f"Лист «{name}»: нет колонок {', '.join(missing)}")
        return []
    unknown = [h for h in header if h and h not in columns]
    if unknown:
        warnings.append(f"Лист «{name}»: колонки {', '.join(unknown)} не используются сайтом")
    out = []
    for number, row in enumerate(rows[1:], start=2):
        record = {h: text(v) for h, v in zip(header, row) if h}
        if any(record.values()):
            out.append((number, record))
    return out


def flag(value: str, where: str, default: bool = True) -> bool:
    v = value.lower()
    if not v:
        return default
    if v in YES:
        return True
    if v in NO:
        return False
    errors.append(f"{where}: «Показывать» должно быть «да» или «нет», а не «{value}»")
    return default


def number(value: str, where: str, column: str, integer: bool = True):
    if not value:
        return None
    v = value.replace(" ", "").replace(",", ".").replace("₽", "").replace("р", "")
    try:
        n = float(v)
    except ValueError:
        errors.append(f"{where}: в «{column}» должно быть число, а не «{value}»")
        return None
    if n < 0:
        errors.append(f"{where}: «{column}» не может быть отрицательным")
        return None
    return int(round(n)) if integer else n


UNITS = {"г": 1, "кг": 1000, "мл": 1, "л": 1000}
GRANULE_ORDER = ["S", "M", "L", "XS"]


def weight_key(label: str):
    """«400 г» → 400, «3,5 кг» → 3500; None, если вес не распознан («1 шт», «1 шт, ~40 г»)."""
    m = re.fullmatch(r"(\d+(?:[.,]\d+)?)\s*(г|кг|мл|л)", label.strip())
    return float(m.group(1).replace(",", ".")) * UNITS[m.group(2)] if m else None


def sort_sizes(items: list[dict]) -> list[dict]:
    """Фасовки по грануле, внутри — по весу. Если вес не везде распознан, порядок из таблицы сохраняется."""
    if any(weight_key(s["w"]) is None for s in items):
        return items
    return sorted(items, key=lambda s: (GRANULE_ORDER.index(s["g"]) if s.get("g") else -1, weight_key(s["w"])))


def percent(value: str) -> str:
    return f"{value.replace('.', ',').rstrip('%')}%" if value else ""


def main() -> None:
    if not WORKBOOK.exists():
        sys.exit(f"Нет файла {WORKBOOK.relative_to(ROOT)}")
    wb = load_workbook(WORKBOOK, data_only=True)
    maker = {p["id"]: p for p in json.loads(MAKER.read_text(encoding="utf-8"))["products"]} if MAKER.exists() else {}

    # ——— марки ———
    brands = []
    for n, r in read_sheet(wb, "Марки", BRAND_COLUMNS):
        where = f"Марки, строка {n}"
        if not r.get("ID") or not r.get("Название"):
            errors.append(f"{where}: нужны ID и Название"); continue
        brands.append({
            "id": r["ID"], "name": r["Название"], "ru": r.get("Подпись", ""), "hue": r.get("Цвет") or "#5CC8F2",
            "cls": r.get("Класс", ""), "badge": r.get("Метка", ""), "about": r.get("Описание", ""),
            "photo": r.get("Фото", ""), "block": flag(r.get("В блоке «Марки»", ""), where),
            "order": number(r.get("Порядок", ""), where, "Порядок") or 9999,
        })
    brand_ids = [b["id"] for b in brands]
    for dup in {i for i in brand_ids if brand_ids.count(i) > 1}:
        errors.append(f"Марки: ID «{dup}» встречается несколько раз")
    hue = {b["id"]: b["hue"] for b in brands}

    # ——— фасовки ———
    sizes: dict[str, list] = {}
    seen_sizes = set()
    for n, r in read_sheet(wb, "Фасовки", SIZE_COLUMNS):
        where = f"Фасовки, строка {n}"
        pid, w = r.get("ID товара", ""), r.get("Фасовка", "")
        if not pid or not w:
            errors.append(f"{where}: нужны «ID товара» и «Фасовка»"); continue
        if not flag(r.get("Показывать", ""), where):
            continue
        g_text = r.get("Гранула", "").lower()
        if g_text and g_text not in GRANULES:
            errors.append(f"{where}: гранула «{r['Гранула']}» — допустимо: {', '.join(GRANULES)}"); continue
        g = GRANULES.get(g_text, "")
        if (pid, g, w) in seen_sizes:
            errors.append(f"{where}: фасовка «{w}»{' / ' + g_text if g_text else ''} у товара {pid} уже есть выше"); continue
        seen_sizes.add((pid, g, w))
        p = number(r.get("Цена, ₽", ""), where, "Цена, ₽")
        p_to = number(r.get("Цена до, ₽", ""), where, "Цена до, ₽")
        if p_to is not None and (p is None or p_to <= p):
            errors.append(f"{where}: «Цена до» должна быть больше «Цены»"); p_to = None
        stock = number(r.get("Остаток", ""), where, "Остаток")
        source = r.get("Источник цены", "")
        if source and source not in PRICE_SOURCES:
            warnings.append(f"{where}: необычный «Источник цены» — «{source}»")
        entry = {"w": w, "p": p}
        if g: entry["g"] = g
        if p_to: entry["pTo"] = p_to
        if stock is not None: entry["stock"] = stock
        sizes.setdefault(pid, []).append(entry)

    # ——— товары ———
    products, product_ids = [], set()
    for n, r in read_sheet(wb, "Товары", PRODUCT_COLUMNS):
        where = f"Товары, строка {n}"
        pid = r.get("ID", "")
        missing = [c for c in REQUIRED["Товары"] if not r.get(c)]
        if missing:
            errors.append(f"{where}: не заполнено {', '.join(missing)}"); continue
        if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", pid):
            errors.append(f"{where}: ID «{pid}» — только латиница в нижнем регистре, цифры и дефис"); continue
        if pid in product_ids:
            errors.append(f"{where}: ID «{pid}» уже есть выше"); continue
        product_ids.add(pid)
        if not flag(r.get("Показывать", ""), where):
            continue
        brand = r["Марка"]
        if brand not in brand_ids:
            errors.append(f"{where}: марки «{brand}» нет на листе «Марки»"); continue
        pet = PETS.get(r["Питомец"].lower())
        if not pet:
            errors.append(f"{where}: питомец «{r['Питомец']}» — допустимо: {', '.join(PETS)}"); continue
        section = r["Раздел"]
        if section not in SECTIONS:
            errors.append(f"{where}: раздел «{section}» — допустимо: {', '.join(SECTIONS)}"); continue
        if pid not in sizes:
            warnings.append(f"{where}: у товара {pid} нет видимых фасовок — на сайт не попадёт"); continue

        m = maker.get(r.get("Карточка производителя", ""), {})
        if r.get("Карточка производителя") and not m:
            warnings.append(f"{where}: карточка производителя «{r['Карточка производителя']}» не найдена")
        an = dict(m.get("an", {}))
        if r.get("Белок, %"): an["Белок"] = percent(r["Белок, %"])
        if r.get("Жир, %"): an["Жир"] = percent(r["Жир, %"])
        if an:  # белок и жир — первыми, как на упаковке
            an = {k: an[k] for k in ["Белок", "Жир"] if k in an} | {k: v for k, v in an.items() if k not in ("Белок", "Жир")}
        tags = [t.strip() for t in r.get("Особенности", "").split(",") if t.strip()] or m.get("tags", [])
        is_food = section in ("Корм сухой", "Корм влажный")
        products.append({
            "id": pid, "brand": brand, "pet": pet, "line": r.get("Линейка") or section, "name": r["Название"],
            "sub": r.get("Коротко") or m.get("sub", ""), "cls": r.get("Класс") or m.get("cls", ""),
            "grain": m.get("grain", ""), "tags": tags, "hue": hue.get(brand, "#5CC8F2"),
            "desc": r.get("Описание") or m.get("desc", ""), "comp": r.get("Состав") or m.get("comp", ""), "an": an,
            "norm": m.get("norm", "") or ("Норма кормления указана на упаковке." if is_food else ""),
            "sizes": sort_sizes(sizes[pid]), "photo": r.get("Фото") or m.get("photo", ""), "normPhoto": m.get("normPhoto", ""),
            "source": m.get("source", ""), "storage": m.get("storage", ""),
            "feedType": m.get("feedType") or FEED_TYPE[section], "section": section,
            **({"type": "treat"} if section == "Лакомства" else {}),
            "_order": number(r.get("Порядок", ""), where, "Порядок") or 99999, "_row": n,
        })
        if any(s["p"] is None for s in sizes[pid]):
            warnings.append(f"{where}: у {pid} есть фасовки без цены — на сайте «цена по запросу»")

    for pid in sizes:
        if pid not in product_ids:
            errors.append(f"Фасовки: товара «{pid}» нет на листе «Товары»")
    for p in products:
        if p["photo"] and not p["photo"].startswith("http") and not (ROOT / p["photo"]).exists():
            warnings.append(f"Товары, строка {p['_row']}: файл фото «{p['photo']}» не найден")

    for w in warnings:
        print("внимание:", w)
    if errors:
        print("\nОшибки — каталог не собран:", *errors, sep="\n  ", file=sys.stderr)
        sys.exit(1)

    products.sort(key=lambda p: (p["_order"], p["_row"]))
    for p in products:
        del p["_order"], p["_row"]
    brands.sort(key=lambda b: b["order"])
    for b in brands:
        del b["order"]
    version = hashlib.sha1(json.dumps([brands, products], ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:10]
    payload = {"generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"), "version": version, "brands": brands, "products": products}
    OUTPUT.write_text("/* Каталог витрины — собирается scripts/import_catalog.py из catalog/Каталог ХаскоШоп.xlsx. Не править вручную. */\n"
                      f"window.XS_CATALOG = {json.dumps(payload, ensure_ascii=False, indent=1)};\n", encoding="utf-8")
    page = PAGE.read_text(encoding="utf-8")
    page_new = re.sub(r'src="assets/data/catalog\.js(\?v=[^"]*)?"', f'src="assets/data/catalog.js?v={version}"', page)
    if page_new != page:
        PAGE.write_text(page_new, encoding="utf-8")
    print(f"готово: товаров — {len(products)}, фасовок — {sum(len(p['sizes']) for p in products)}, марок — {len(brands)} → {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
